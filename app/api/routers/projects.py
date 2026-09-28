import csv
import io
from datetime import date
from typing import Literal
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from openpyxl import load_workbook
from app.core.database import get_db
from app.core.deps import require_roles
from app.core.permissions import ADMIN_ROLES, READ_ROLES
from app.core.audit import record_audit
from app.models.identity import User
from app.models.projects import Project, ProjectGroup, ProjectExpectedGroup
from app.models.whatsapp import WhatsAppGroup
from app.models.applications import Application
from app.services.projects import business_key, project_dashboard
router = APIRouter(prefix='/projects', tags=['projects'])

class NameIn(BaseModel):
    name: str = Field(min_length=1, max_length=150)
class GroupsIn(BaseModel):
    group_ids: list[str] = Field(max_length=1000)

def require_project(db, project_id):
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(404, 'Projet introuvable')
    return project

@router.get('')
def list_projects(db: Session = Depends(get_db), user: User = Depends(require_roles(*READ_ROLES))):
    return [{'id': p.id, 'name': p.name} for p in db.query(Project).order_by(Project.name)]

@router.post('')
def create_project(payload: NameIn, db: Session = Depends(get_db), user: User = Depends(require_roles(*ADMIN_ROLES))):
    name = payload.name.strip()
    if not name or db.query(Project).filter_by(name=name).first():
        raise HTTPException(409, 'Nom vide ou projet déjà existant')
    project = Project(name=name)
    db.add(project)
    db.flush()
    record_audit(db, user_id=user.id, action='project.create', entity_id=project.id)
    db.commit()
    return {'id': project.id, 'name': project.name}

@router.get('/{project_id}/groups')
def groups(project_id: str, db: Session = Depends(get_db), user: User = Depends(require_roles(*READ_ROLES))):
    require_project(db, project_id)
    links = {r.group_id: r.project_id for r in db.query(ProjectGroup)}
    return [{'id': g.id, 'name': g.name, 'external_id': g.whatsapp_group_external_id, 'project_id': links.get(g.id)} for g in db.query(WhatsAppGroup).order_by(WhatsAppGroup.name)]

@router.put('/{project_id}/groups')
def set_groups(project_id: str, payload: GroupsIn, db: Session = Depends(get_db), user: User = Depends(require_roles(*ADMIN_ROLES))):
    require_project(db, project_id)
    ids = set(payload.group_ids)
    if db.query(WhatsAppGroup).filter(WhatsAppGroup.id.in_(ids)).count() != len(ids):
        raise HTTPException(422, 'Groupe introuvable')
    if db.query(ProjectGroup).filter(ProjectGroup.group_id.in_(ids), ProjectGroup.project_id != project_id).first():
        raise HTTPException(409, 'Un groupe est déjà rattaché à un autre projet ; retirez son rattachement d’abord')
    db.query(ProjectGroup).filter_by(project_id=project_id).delete()
    db.add_all([ProjectGroup(group_id=g, project_id=project_id) for g in ids])
    record_audit(db, user_id=user.id, action='project.groups', entity_id=project_id, details={'count': len(ids)})
    db.commit()
    return {'groups': len(ids)}

@router.post('/{project_id}/expected')
async def import_expected(project_id: str, application_id: str = Form(...), file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(require_roles(*ADMIN_ROLES))):
    require_project(db, project_id)
    if not db.get(Application, application_id):
        raise HTTPException(422, 'Application introuvable')
    content = await file.read(2_000_001)
    if len(content) > 2_000_000:
        raise HTTPException(413, 'Fichier limité à 2 Mo')
    try:
        if (file.filename or '').lower().endswith('.xlsx'):
            with io.BytesIO(content) as stream:
                book = load_workbook(stream, read_only=True, data_only=True)
                try:
                    values = []
                    for i, row in enumerate(book.active.iter_rows(values_only=True)):
                        if i > 5000:
                            raise ValueError('Maximum 5000 lignes')
                        values.append(row[0] if row else None)
                finally:
                    book.close()
        else:
            text = content.decode('utf-8-sig')
            values = [row[0] for row in csv.reader(io.StringIO(text), delimiter=';' if ';' in text.splitlines()[0] else ',') if row]
        keys = {business_key(v) for v in values if v is not None and business_key(v)}
        keys -= {'id', 'group_id', 'id_groupe', 'identifiant'}
        if not keys or len(keys) > 5000 or any(len(k) > 200 for k in keys):
            raise ValueError('Liste vide, ID trop long ou plus de 5000 ID')
    except Exception as exc:
        raise HTTPException(422, 'Fichier invalide : première colonne = ID métier, CSV UTF-8 ou Excel, maximum 5000 ID') from exc
    existing = {r.business_id for r in db.query(ProjectExpectedGroup).filter_by(project_id=project_id, application_id=application_id)}
    db.add_all([ProjectExpectedGroup(project_id=project_id, application_id=application_id, business_id=k) for k in keys - existing])
    record_audit(db, user_id=user.id, action='project.import_ids', entity_id=project_id, details={'added': len(keys-existing), 'application_id': application_id})
    db.commit()
    return {'added': len(keys-existing), 'already_present': len(keys & existing)}

@router.get('/{project_id}/dashboard')
def dashboard(project_id: str, start: date, end: date, cadence: Literal['daily', 'weekly', 'period'] = 'period', application_id: str | None = None, db: Session = Depends(get_db), user: User = Depends(require_roles(*READ_ROLES))):
    require_project(db, project_id)
    if end < start or (end-start).days > 92:
        raise HTTPException(422, 'Période de 1 à 93 jours')
    return project_dashboard(db, project_id, start, end, cadence, application_id)
