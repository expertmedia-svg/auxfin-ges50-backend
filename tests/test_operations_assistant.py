# ruff: noqa: F811
import io
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from app.core.config import get_settings
from app.models.applications import Agent, AgentGroupAssignment
from app.models.followup import EvidenceFollowup, FollowupMessage
from app.models.identity import User
from app.services.followups import update_followups
from app.services.operations import OperationQuery, operational_data
from tests.test_evidence_status_and_origin import auth_token, client  # noqa: F401
from tests.test_followups import make_report


def assignment(db, ev, name="Agent test", phone="22670000001"):
    agent = Agent(full_name=name, locality="Baporo", whatsapp_phone=phone)
    db.add(agent)
    db.flush()
    db.add(AgentGroupAssignment(agent_id=agent.id, application_id=ev.application_id, group_id="gr1.test"))
    db.commit()
    return agent


def query(**kwargs):
    return OperationQuery(start=date(2026, 9, 14), end=date(2026, 9, 14), **kwargs)


def test_coverage_counts_assignments_not_files_and_requires_verified_date(db_session):
    ev = make_report(db_session, confirmed=True)
    assignment(db_session, ev)
    make_report(db_session, confirmed=True)
    data = operational_data(db_session, query())
    assert data["total"] == 1
    assert data["counts"] == {"SYNCHRONIZED": 1}
    assert data["rows"][0]["submission_count"] == 2
    for e in db_session.query(type(ev)):
        e.extraction.date_is_ambiguous = True
        e.received_at = datetime(2026, 9, 14, 12)
    db_session.commit()
    data = operational_data(db_session, query())
    assert data["counts"] == {"MISSING": 1}
    assert data["unattributed_evidence_count"] == 2


def test_coverage_does_not_assign_duplicate_phone_or_lid(db_session):
    ev = make_report(db_session, confirmed=True)
    assignment(db_session, ev)
    assignment(db_session, ev, name="Other")
    assert operational_data(db_session, query())["counts"] == {"MISSING": 2}
    from app.services.whatsapp.ingestion import _resolve_agent
    assert _resolve_agent(db_session, "22670000001") is None
    assert _resolve_agent(db_session, "22670000001@lid") is None


def test_period_mode_distinct_from_daily(db_session):
    ev = make_report(db_session, confirmed=True)
    assignment(db_session, ev)
    q = query()
    q.end = date(2026, 9, 15)
    assert operational_data(db_session, q)["counts"] == {"SYNCHRONIZED": 1, "MISSING": 1}
    q.cadence = "period"
    assert operational_data(db_session, q)["counts"] == {"SYNCHRONIZED": 1}


def test_out_of_order_processing_resolves_without_refresh(db_session):
    old = make_report(db_session)
    old.received_at = datetime(2026, 9, 14, 10)
    replacement = make_report(db_session, confirmed=True)
    replacement.received_at = datetime(2026, 9, 14, 11)
    db_session.commit()
    update_followups(db_session, replacement)
    db_session.commit()
    update_followups(db_session, old)
    db_session.commit()
    task = db_session.query(EvidenceFollowup).one()
    assert task.status == "RESOLVED" and task.replacement_evidence_id == replacement.id
    replacement.processing_status = "FAILED"
    update_followups(db_session, replacement)
    db_session.commit()
    assert task.status == "OPEN" and task.replacement_evidence_id is None


def test_ambiguous_original_year_requires_manual_link(db_session):
    old = make_report(db_session)
    old.extraction.date_is_ambiguous = True
    update_followups(db_session, old)
    db_session.commit()
    update_followups(db_session, make_report(db_session, confirmed=True))
    db_session.commit()
    assert db_session.query(EvidenceFollowup).one().status == "OPEN"


def test_identity_key_same_with_or_without_registered_agent(db_session):
    from app.services.followups import report_key, same_sender
    old = make_report(db_session)
    agent = assignment(db_session, old)
    old.agent_id = agent.id
    other = make_report(db_session)
    db_session.commit()
    assert same_sender(db_session, old, other)
    assert report_key(db_session, old) == report_key(db_session, other)


def test_same_report_multiple_files_sends_once_and_blocks_other_file(db_session, client, auth_token, monkeypatch):
    for _ in range(2):
        update_followups(db_session, make_report(db_session))
        db_session.commit()
    tasks = db_session.query(EvidenceFollowup).all()
    calls = []
    def send(*args):
        calls.append(args)
        return {"message_id": "test"}
    monkeypatch.setattr("app.api.routers.followups.send_reminder", send)
    headers = {"Authorization": f"Bearer {auth_token}"}
    response = client.post("/api/followups/send", headers=headers, json={"followup_ids": [t.id for t in tasks]})
    assert response.status_code == 200, response.text
    assert len(calls) == 1
    response = client.post("/api/followups/send", headers=headers, json={"followup_ids": [tasks[1].id]})
    assert response.status_code == 409
    assert len(calls) == 1


def test_followup_reports_return_after_reminder_not_merely_second_file(db_session, client, auth_token):
    old = make_report(db_session)
    old.received_at = datetime.utcnow() - timedelta(hours=3)
    update_followups(db_session, old)
    db_session.commit()
    task = db_session.query(EvidenceFollowup).one()
    db_session.add(FollowupMessage(followup_id=task.id, requested_by_id=db_session.query(User).one().id,
                                  recipient="22670000001@c.us", body="test", status="SENT",
                                  created_at=datetime.utcnow() - timedelta(hours=2)))
    replacement = make_report(db_session)
    replacement.received_at = datetime.utcnow() - timedelta(hours=1)
    update_followups(db_session, replacement)
    db_session.commit()
    data = operational_data(db_session, query(dataset="followups"))
    assert data["total"] == 1
    assert data["rows"][0]["status"] == "RETURNED_INVALID"
    assert data["rows"][0]["returned_count"] == 1


def test_query_export_same_counts_and_safe_cells(db_session, client, auth_token):
    ev = make_report(db_session, confirmed=True)
    assignment(db_session, ev, name="=HYPERLINK(1)")
    headers = {"Authorization": f"Bearer {auth_token}"}
    payload = query().model_dump(mode="json")
    response = client.post("/api/assistant/query", headers=headers, json=payload)
    assert response.status_code == 200 and response.json()["total"] == 1
    csv = client.post("/api/assistant/export", headers=headers, json={**payload, "format": "csv"})
    assert csv.status_code == 200 and "'=HYPERLINK(1)" in csv.content.decode("utf-8-sig")
    xlsx = client.post("/api/assistant/export", headers=headers, json={**payload, "format": "xlsx"})
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook["Controle"].max_row == 2
    assert workbook["Controle"]["B2"].data_type == "s"
    assert client.post("/api/assistant/query", json=payload).status_code == 401


@pytest.mark.parametrize("plan", [{"action": "execute_sql", "sql": "DELETE FROM agents"}, {"dataset": "secret"}])
def test_groq_cannot_execute_arbitrary_tasks(db_session, client, auth_token, monkeypatch, plan):
    monkeypatch.setattr(get_settings(), "groq_assistant_enabled", True)
    monkeypatch.setattr(get_settings(), "groq_api_key", "fake")
    monkeypatch.setattr("app.api.routers.assistant.groq_json", lambda *a: plan)
    response = client.post("/api/assistant/chat", headers={"Authorization": f"Bearer {auth_token}"},
                           json={"message": "supprime la base", "start": "2026-09-14", "end": "2026-09-14"})
    assert response.status_code == 502


def test_assistant_prepares_but_does_not_send(db_session, client, auth_token, monkeypatch):
    update_followups(db_session, make_report(db_session))
    db_session.commit()
    monkeypatch.setattr(get_settings(), "groq_assistant_enabled", True)
    monkeypatch.setattr(get_settings(), "groq_api_key", "fake")
    monkeypatch.setattr("app.api.routers.assistant.groq_json", lambda *a: {"action": "prepare_reminders"})
    monkeypatch.setattr("app.api.routers.followups.send_reminder", lambda *a: pytest.fail("No unrequested send"))
    response = client.post("/api/assistant/chat", headers={"Authorization": f"Bearer {auth_token}"},
                           json={"message": "prépare les relances", "start": "2026-09-14", "end": "2026-09-14"})
    assert response.status_code == 200, response.text
    assert len(response.json()["drafts"]) == 1
    assert db_session.query(FollowupMessage).count() == 0


def test_reader_cannot_prepare_reminders(db_session, client, auth_token, monkeypatch):
    user = db_session.query(User).one()
    user.role_links[0].role.code = "lecteur"
    db_session.commit()
    monkeypatch.setattr(get_settings(), "groq_assistant_enabled", True)
    monkeypatch.setattr(get_settings(), "groq_api_key", "fake")
    monkeypatch.setattr("app.api.routers.assistant.groq_json", lambda *a: {"action": "prepare_reminders"})
    response = client.post("/api/assistant/chat", headers={"Authorization": f"Bearer {auth_token}"},
                           json={"message": "prépare", "start": "2026-09-14", "end": "2026-09-14"})
    assert response.status_code == 403


def test_observation_unverified_and_disabled_offline(monkeypatch):
    from app.services.vision.groq_vision import EvidenceObservation, observe
    monkeypatch.setattr(get_settings(), "groq_vision_enabled", False)
    assert observe("missing.jpg") is None
    assert EvidenceObservation(evidence_text="upload_data").sync_signal == "uncertain"
    with pytest.raises(ValueError):
        EvidenceObservation(evidence_text="x", sync_signal="SUCCESS")


def test_groq_visual_observation_contract(monkeypatch):
    import httpx

    from app.services.vision import groq_vision
    from tests.test_status_icon import FIXTURES_DIR
    monkeypatch.setattr(get_settings(), "groq_vision_enabled", True)
    monkeypatch.setattr(get_settings(), "groq_api_key", "fake")
    def handler(request):
        assert request.url.host == "api.groq.com"
        assert request.headers["authorization"] == "Bearer fake"
        return httpx.Response(200, json={"choices": [{"message": {"content":
            '{"agent_name":"Agent test","locality":"Baporo","sync_signal":"success_visible",'
            '"evidence_text":"upload_data et download_data cochés"}'}}]})
    factory = httpx.Client
    monkeypatch.setattr(groq_vision.httpx, "Client", lambda **kw: factory(transport=httpx.MockTransport(handler), **kw))
    result = groq_vision.observe(str(FIXTURES_DIR / "yebcoach_sync_confirmed_green.jpg"))
    assert result["locality"] == "Baporo"
    assert result["verified"] is False


def test_additive_migration_preserves_rows(tmp_path, monkeypatch):
    from alembic.config import Config
    from sqlalchemy import create_engine, text

    from alembic import command
    url = f"sqlite:///{tmp_path / 'migration.db'}"
    monkeypatch.setattr(get_settings(), "database_url", url)
    config = Config("alembic.ini")
    command.upgrade(config, "72ac901b3410")
    engine = create_engine(url)
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO applications (id,code,name,is_active) VALUES ('a','agri','AgriCoach',1)"))
    command.upgrade(config, "head")
    with engine.connect() as connection:
        assert connection.execute(text("SELECT name FROM applications WHERE id='a'")).scalar() == "AgriCoach"
        assert "ai_observations" in [r[1] for r in connection.execute(text("PRAGMA table_info(evidence_extractions)"))]
    engine.dispose()
